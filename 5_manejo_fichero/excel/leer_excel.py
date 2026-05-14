from openpyxl import load_workbook

# Leer de un archivo Excel
excel_file = load_workbook('./data/empleados.xlsx')

# Seleccionar la hoja activa del archivo Excel
hoja = excel_file.active   

# Recorrar filas y columnas de la hoja excel
for fila in hoja.iter_rows(min_row=2, values_only=True):
    for value in fila:
        print(value)
