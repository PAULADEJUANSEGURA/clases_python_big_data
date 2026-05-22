import os
from openpyxl import load_workbook, Workbook
from lib.functions import limpiar_id, limpiar_texto, limpiar_precio, limpiar_stock  

def cargar_excel(carpeta, archivo):
   excel = load_workbook(f"./{carpeta}/{archivo}")
   hoja = excel.active
   hoja = excel['Inventario']
   
   #cargamos las cabeceras
   filas = hoja.iter_rows(values_only = True)
   cabeceras = next(filas)
   
   lista_resultante = []
   for fila in hoja.iter_rows(min_row=2, values_only=True):
       producto = dict(zip(cabeceras, fila))
       lista_resultante.append(producto)
       
   return lista_resultante


def crear_excel(carpeta, archivo, datos):
    pass

def crear_excel(carpeta, archivo, datos):
    os.makedirs(carpeta, exist_ok=True)
    ruta = f"./{carpeta}/{archivo}"
    # Creamos el libro de excel
    excel = Workbook()
    # Creamos la hoja de excel
    hoja = excel.active
    hoja.title = "Inventario_limpio"
    # Escribimos las cabeceras
    cabeceras = list(datos[0].keys())
    hoja.append(cabeceras)
    # Escribimos los datos
    for producto in datos:
        fila = [producto[clave] for clave in cabeceras]
        hoja.append(fila)

    excel.save(ruta)
    print(f"Archivo '{archivo}' creado exitosamente en la carpeta '{carpeta}'.")

def procesar(datos):
    lista_limpia_vacia = []
    for item in datos:
        item_limpio = {
            'id_producto': limpiar_id(item['id_producto']),
            'nombre_producto': limpiar_texto(item['nombre_producto']),
            'categoria': limpiar_texto(item['categoria']),
            'precio': limpiar_precio(item['precio']),
            'stock': limpiar_stock(item['stock']),
        }
    lista_limpia_vacia.append(item_limpio)

datos = cargar_excel("data", "inventario_sucio.xlsx")
datos_limpios = procesar(datos)
crear_excel("data", "inventario_limpio.xlsx", datos_limpios )